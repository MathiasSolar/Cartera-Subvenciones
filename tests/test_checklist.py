from io import BytesIO

import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("CARTERA_DB", str(tmp_path / "test.db"))
    from app.main import app

    with TestClient(app) as c:
        yield c


def proyecto(client, **extra):
    r = client.post("/api/proyectos", json={"nombre": "Con checklist", "codigo": "8%-CUL-014", **extra})
    assert r.status_code == 201, r.text
    return r.json()


def marcar(client, pid, clave, hecho=True):
    return client.put(f"/api/proyectos/{pid}/checklist/{clave}", json={"hecho": hecho})


def claves_de(client, etapa):
    from app.checklist import claves

    return [c for t in client.get("/api/config").json()["checklist_por_etapa"][etapa] for c in claves(t)]


def completar(client, pid, etapa):
    for clave in claves_de(client, etapa):
        assert marcar(client, pid, clave).status_code == 200


def test_adjudicado_parte_con_el_checklist_vacio(client):
    cfg = client.get("/api/config").json()["checklist_por_etapa"]
    assert [t["clave"] for t in cfg["Adjudicado"]] == ["secpir", "carpeta", "docdigital", "firmado"]
    assert len(cfg["Adjudicado"][1]["subtareas"]) == 7   # CDP, resolución, declaración jurada…
    assert "{codigo}" in " ".join(cfg["Adjudicado"][2]["instrucciones"])

    p = proyecto(client)
    assert p["checklist"] == {"etapa": "Adjudicado", "hechas": {}, "total": 4, "listas": 0, "completo": False}


def test_marcar_queda_en_la_bitacora_y_las_subtareas_completan_la_tarea(client):
    p = proyecto(client)
    r = marcar(client, p["id"], "secpir")
    assert r.status_code == 200
    assert r.json()["bitacora"][0]["texto"] == "✓ Llenar la información en SECPIR"
    assert r.json()["checklist"]["listas"] == 1 and "secpir" in r.json()["checklist"]["hechas"]
    assert marcar(client, p["id"], "secpir").json()["checklist"]["listas"] == 1   # marcar dos veces no duplica

    r = marcar(client, p["id"], "doc-cdp")
    assert r.json()["bitacora"][0]["texto"] == "✓ Juntar los documentos en la carpeta del proyecto: CDP"
    assert r.json()["checklist"]["listas"] == 1        # la carpeta aún no: faltan 6 documentos
    for clave in claves_de(client, "Adjudicado")[2:8]:
        r = marcar(client, p["id"], clave)
    assert r.json()["checklist"]["listas"] == 2

    r = marcar(client, p["id"], "doc-core", hecho=False)
    assert r.json()["bitacora"][0]["texto"] == "Se desmarcó: Juntar los documentos en la carpeta del proyecto: Acuerdo CORE"
    assert r.json()["checklist"]["listas"] == 1

    assert marcar(client, p["id"], "carpeta").status_code == 404     # tiene subtareas: se marca cada una
    assert marcar(client, p["id"], "pagare").status_code == 404      # es del checklist de Convenio
    assert marcar(client, 999, "secpir").status_code == 404


def test_la_proxima_accion_sigue_al_checklist(client):
    p = proyecto(client, accion="Llenar la información en SECPIR")
    r = marcar(client, p["id"], "secpir")
    assert r.json()["accion"] == "Juntar los documentos en la carpeta del proyecto"
    # Si la próxima acción es otra cosa, no se toca
    r = marcar(client, p["id"], "docdigital")
    assert r.json()["accion"] == "Juntar los documentos en la carpeta del proyecto"


def test_pasar_a_convenio_exige_el_checklist_completo(client):
    p = proyecto(client)
    client.post(f"/api/proyectos/{p['id']}/bitacora", json={"texto": "Una nota no basta"})
    marcar(client, p["id"], "secpir")
    r = client.post(f"/api/proyectos/{p['id']}/avanzar")
    assert r.status_code == 409
    detalle = r.json()["detail"]
    assert "checklist de Adjudicado" in detalle and "Subir el CDP" in detalle and "SECPIR" not in detalle

    completar(client, p["id"], "Adjudicado")
    r = client.post(f"/api/proyectos/{p['id']}/avanzar")
    assert r.status_code == 200
    assert r.json()["etapa"] == "Convenio" and r.json()["accion"] == "Generar el convenio"
    assert r.json()["checklist"] == {"etapa": "Convenio", "hechas": r.json()["checklist"]["hechas"],
                                     "total": 4, "listas": 0, "completo": False}

    assert client.post(f"/api/proyectos/{p['id']}/avanzar").status_code == 409
    completar(client, p["id"], "Convenio")
    r = client.post(f"/api/proyectos/{p['id']}/avanzar")
    assert r.status_code == 200 and r.json()["etapa"] == "Transferencia"


def test_no_se_puede_saltar_un_checklist(client):
    p = proyecto(client)
    completar(client, p["id"], "Adjudicado")
    datos = {k: p[k] for k in ["nombre", "codigo", "linea", "organizacion", "monto", "contacto", "accion", "fecha", "notas"]}
    r = client.put(f"/api/proyectos/{p['id']}", json={**datos, "etapa": "Transferencia"})
    assert r.status_code == 409 and "hay que pasar por Convenio" in r.json()["detail"]
    assert client.put(f"/api/proyectos/{p['id']}", json={**datos, "etapa": "Convenio"}).status_code == 200


def test_pagare_vence_un_anio_despues_de_la_ultima_rendicion(client):
    p = proyecto(client, etapa="Convenio")
    r = client.put(f"/api/proyectos/{p['id']}/pagare", json={"ultima_rendicion": "2027-01-15"})
    assert r.status_code == 200
    assert r.json()["ultima_rendicion"] == "2027-01-15" and r.json()["vence_pagare"] == "2028-01-15"
    assert r.json()["bitacora"][0]["texto"] == "Pagaré: vence el 15-01-2028 (última rendición: 15-01-2027)"

    # Se puede corregir el vencimiento a mano
    r = client.put(f"/api/proyectos/{p['id']}/pagare", json={"ultima_rendicion": "2027-01-15", "vence_pagare": "2028-02-01"})
    assert r.json()["vence_pagare"] == "2028-02-01"
    # 29 de febrero → 28 de febrero del año siguiente
    r = client.put(f"/api/proyectos/{p['id']}/pagare", json={"ultima_rendicion": "2028-02-29"})
    assert r.json()["vence_pagare"] == "2029-02-28"

    r = client.put(f"/api/proyectos/{p['id']}/pagare", json={})
    assert r.json()["vence_pagare"] is None
    assert r.json()["bitacora"][0]["texto"] == "Pagaré: se quitó la fecha de vencimiento"


def test_excel_trae_el_vencimiento_del_pagare(client):
    from datetime import datetime
    from openpyxl import load_workbook

    p = proyecto(client, etapa="Convenio")
    client.put(f"/api/proyectos/{p['id']}/pagare", json={"ultima_rendicion": "2027-01-15"})
    wb = load_workbook(BytesIO(client.get("/api/exportar/rendiciones.xlsx").content))
    filas = [[c.value for c in f] for f in wb["Resumen"].iter_rows()]
    col = filas[0].index("Vence pagaré")
    assert filas[1][col] == datetime(2028, 1, 15)


def test_se_pueden_corregir_tareas_de_etapas_anteriores(client):
    p = proyecto(client, etapa="Transferencia")
    assert marcar(client, p["id"], "secpir").status_code == 200
    r = marcar(client, p["id"], "secpir", hecho=False)
    assert r.status_code == 200 and "secpir" not in r.json()["checklist_hechas"]
    r = marcar(client, p["id"], "pagare")
    assert r.status_code == 200 and "pagare" in r.json()["checklist_hechas"]
    assert r.json()["checklist"] is None   # Transferencia no tiene checklist propio

    p2 = proyecto(client, nombre="Recién adjudicado")
    assert marcar(client, p2["id"], "convenio").status_code == 404   # etapa futura


def test_migracion_completa_los_checklists_de_etapas_ya_pasadas(tmp_path):
    import sqlite3
    from app import db

    ruta = tmp_path / "vieja.db"
    c = sqlite3.connect(ruta)
    for n, sql in enumerate(db.MIGRACIONES[:8], start=1):
        c.executescript(sql)
        c.execute(f"PRAGMA user_version = {n}")
    t = "2026-09-01T12:00:00+00:00"
    for nombre, etapa in [("A", "Adjudicado"), ("C", "Convenio"), ("T", "Transferencia"), ("R", "Rendición")]:
        c.execute("INSERT INTO proyectos (nombre, etapa, creado, actualizado) VALUES (?, ?, ?, ?)", (nombre, etapa, t, t))
    ids = {r[1]: r[0] for r in c.execute("SELECT id, nombre FROM proyectos")}
    c.execute("INSERT INTO checklist (proyecto_id, clave, hecho) VALUES (?, 'secpir', ?)", (ids["A"], t))
    c.execute("INSERT INTO checklist (proyecto_id, clave, hecho) VALUES (?, 'convenio', ?)", (ids["C"], t))
    c.commit(); c.close()

    from app.checklist import claves, tareas

    adj = {k for tarea in tareas("Adjudicado") for k in claves(tarea)}
    conv = {k for tarea in tareas("Convenio") for k in claves(tarea)}
    c = db.conectar(ruta)
    marcas = lambda nombre: {r[0] for r in c.execute("SELECT clave FROM checklist WHERE proyecto_id = ?", (ids[nombre],))}  # noqa: E731
    assert marcas("A") == {"secpir"}                          # sigue en Adjudicado: no se toca
    assert marcas("C") == adj | {"convenio"}                  # Adjudicado completo; Convenio queda como estaba
    assert marcas("T") == adj | conv and marcas("R") == adj | conv
    notas = [r[0] for r in c.execute("SELECT texto FROM bitacora WHERE proyecto_id = ?", (ids["T"],))]
    assert notas == ["Checklist de Adjudicado marcado como completo: el proyecto ya había pasado esa etapa",
                     "Checklist de Convenio marcado como completo: el proyecto ya había pasado esa etapa"]
    c.close()


def test_migracion_domo_pasa_a_docdigital(tmp_path):
    import sqlite3
    from app import db

    ruta = tmp_path / "vieja.db"
    c = sqlite3.connect(ruta)
    for n, sql in enumerate(db.MIGRACIONES[:9], start=1):
        c.executescript(sql)
        c.execute(f"PRAGMA user_version = {n}")
    t = "2026-10-01T12:00:00+00:00"
    c.execute("INSERT INTO proyectos (nombre, etapa, accion, creado, actualizado) VALUES "
              "('A', 'Adjudicado', 'Subir el CDP y la resolución a DOMO DPIR para la firma de las jefaturas', ?, ?)", (t, t))
    c.execute("INSERT INTO checklist (proyecto_id, clave, hecho) VALUES (1, 'domo', ?)", (t,))
    c.commit(); c.close()

    c = db.conectar(ruta)
    assert c.execute("SELECT accion FROM proyectos").fetchone()[0] == \
        "Subir el CDP y la resolución a DocDigital para la firma de las jefaturas"
    assert [r[0] for r in c.execute("SELECT clave FROM checklist")] == ["docdigital"]
    c.close()


def test_proyecto_nuevo_parte_con_su_primer_paso(client):
    p = client.post("/api/proyectos", json={"nombre": "Sin acción"}).json()
    assert p["accion"] == "Llenar la información en SECPIR" and p["fecha"] and p["accion_manual"] == 0
    p = client.post("/api/proyectos", json={"nombre": "Con acción", "accion": "Llamar al municipio"}).json()
    assert p["accion"] == "Llamar al municipio"


def test_siguiente_paso_automatico_o_manual(client):
    p = proyecto(client, etapa="Ejecución")
    r = client.put(f"/api/proyectos/{p['id']}/siguiente",
                   json={"accion": " Visitar la ejecución ", "fecha": "2026-11-02", "manual": True})
    assert r.status_code == 200
    assert (r.json()["accion"], r.json()["fecha"], r.json()["accion_manual"]) == ("Visitar la ejecución", "2026-11-02", 1)

    # Al avanzar, la nueva etapa vuelve a la próxima acción automática
    client.post(f"/api/proyectos/{p['id']}/bitacora", json={"texto": "Ejecución terminada"})
    r = client.post(f"/api/proyectos/{p['id']}/avanzar")
    assert r.json()["etapa"] == "Rendición" and r.json()["accion_manual"] == 0
    assert r.json()["accion"] == "Definir el período de rendiciones"
    assert client.put("/api/proyectos/999/siguiente", json={"accion": "x"}).status_code == 404
