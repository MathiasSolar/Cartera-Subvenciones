import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("CARTERA_DB", str(tmp_path / "test.db"))
    from app.main import app

    with TestClient(app) as c:
        yield c


def nuevo(client, **extra):
    datos = {"nombre": "Proyecto de prueba", "monto": 1_000_000, "fecha": "2026-10-01", **extra}
    r = client.post("/api/proyectos", json=datos)
    assert r.status_code == 201, r.text
    return r.json()


def test_crear_y_listar(client):
    p = nuevo(client)
    assert p["etapa"] == "Adjudicado"   # todo proyecto parte adjudicado
    assert p["bitacora"][0]["sistema"] == 1

    lista = client.get("/api/proyectos").json()
    assert [x["nombre"] for x in lista] == ["Proyecto de prueba"]


def nota(client, pid, texto="Revisé los antecedentes"):
    assert client.post(f"/api/proyectos/{pid}/bitacora", json={"texto": texto}).status_code == 201


def test_cambio_de_etapa_queda_en_bitacora(client):
    p = nuevo(client, etapa="Ejecución")
    datos = {k: p[k] for k in ["nombre", "codigo", "linea", "organizacion", "monto",
                               "contacto", "accion", "fecha", "notas"]}
    # En una etapa sin checklist, avanzar sin registro en la bitácora no se permite
    r = client.put(f"/api/proyectos/{p['id']}", json={**datos, "etapa": "Rendición"})
    assert r.status_code == 409 and "bitácora" in r.json()["detail"]
    nota(client, p["id"])
    r = client.put(f"/api/proyectos/{p['id']}", json={**datos, "etapa": "Rendición"})
    assert r.status_code == 200
    assert r.json()["bitacora"][0]["texto"] == "Etapa: Ejecución → Rendición"
    # Retroceder para corregir sí se permite sin registro
    assert client.put(f"/api/proyectos/{p['id']}", json={**datos, "etapa": "Ejecución"}).status_code == 200


def test_avanzar_exige_registro_y_pone_el_primer_paso(client):
    p = nuevo(client, etapa="Ejecución", accion="Algo antiguo")
    r = client.post(f"/api/proyectos/{p['id']}/avanzar")
    assert r.status_code == 409 and "Rendición" in r.json()["detail"]

    nota(client, p["id"], "Ejecución terminada")
    r = client.post(f"/api/proyectos/{p['id']}/avanzar")
    assert r.status_code == 200
    assert r.json()["etapa"] == "Rendición"
    assert r.json()["accion"] == "Definir el período de rendiciones" and r.json()["fecha"]

    # En la nueva etapa hay que volver a registrar antes de avanzar
    assert client.post(f"/api/proyectos/{p['id']}/avanzar").status_code == 409


def test_avanzar_hasta_el_final(client):
    p = nuevo(client, etapa="Rendición")
    nota(client, p["id"])
    r = client.post(f"/api/proyectos/{p['id']}/avanzar")
    assert r.json()["etapa"] == "Cerrado" and r.json()["accion"] == ""
    nota(client, p["id"])
    r = client.post(f"/api/proyectos/{p['id']}/avanzar")
    assert r.status_code == 409


def test_nota_y_borrado(client):
    p = nuevo(client)
    r = client.post(f"/api/proyectos/{p['id']}/bitacora", json={"texto": "Llamé a la organización"})
    assert r.json()["bitacora"][0]["texto"] == "Llamé a la organización"

    assert client.delete(f"/api/proyectos/{p['id']}").status_code == 204
    assert client.get(f"/api/proyectos/{p['id']}").status_code == 404


def test_validaciones(client):
    assert client.post("/api/proyectos", json={"nombre": "  "}).status_code == 422
    assert client.post("/api/proyectos", json={"nombre": "X", "etapa": "Inventada"}).status_code == 422
    assert client.post("/api/proyectos", json={"nombre": "X", "monto": -5}).status_code == 422


def test_sirve_la_interfaz(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "Cartera DIPIR" in r.text


def test_tema_se_guarda(client):
    assert client.get("/api/config").json()["tema"] is None

    r = client.put("/api/ajustes/tema", json={"tema": "oscuro"})
    assert r.status_code == 200
    assert client.get("/api/config").json()["tema"] == "oscuro"

    client.put("/api/ajustes/tema", json={"tema": "claro"})
    assert client.get("/api/config").json()["tema"] == "claro"

    assert client.put("/api/ajustes/tema", json={"tema": "rosado"}).status_code == 422


def periodo(client, pid, desde, hasta):
    return client.put(f"/api/proyectos/{pid}/rendiciones/periodo", json={"desde": desde, "hasta": hasta})


def test_periodo_crea_las_rendiciones(client):
    p = nuevo(client, monto=6_000_000)
    r = periodo(client, p["id"], "2026-09", "2027-02")
    assert r.status_code == 200
    meses = [x["mes"] for x in r.json()["rendiciones"]]
    assert meses == ["2026-09", "2026-10", "2026-11", "2026-12", "2027-01", "2027-02"]
    assert all(x["estado"] == "Pendiente" for x in r.json()["rendiciones"])
    assert "sep 2026 a feb 2027 (6 rendiciones)" in r.json()["bitacora"][0]["texto"]

    # Proyecto de un solo mes: una rendición
    q = nuevo(client, nombre="Proyecto corto")
    assert len(periodo(client, q["id"], "2026-09", "2026-09").json()["rendiciones"]) == 1

    # La lista trae los estados para la barra de rendiciones
    lista = {x["id"]: x for x in client.get("/api/proyectos").json()}
    assert len(lista[p["id"]]["rendiciones"]) == 6

    # La lista trae el monto de cada rendición, para mostrar lo rendido bajo el monto total
    client.put(f"/api/rendiciones/{r.json()['rendiciones'][0]['id']}", json={"estado": "Aprobada", "monto": 750_000})
    lista = {x["id"]: x for x in client.get("/api/proyectos").json()}
    assert sum(x["monto"] or 0 for x in lista[p["id"]]["rendiciones"]) == 750_000


def test_no_borra_meses_con_datos(client):
    p = nuevo(client)
    rend = periodo(client, p["id"], "2026-09", "2026-12").json()["rendiciones"]
    diciembre = rend[-1]["id"]
    client.put(f"/api/rendiciones/{diciembre}", json={"estado": "Aprobada", "monto": 500_000})

    r = periodo(client, p["id"], "2026-09", "2026-11")
    assert r.status_code == 409
    assert "dic 2026" in r.json()["detail"]

    # Extender sí se puede y conserva lo ya registrado
    r = periodo(client, p["id"], "2026-08", "2027-01")
    meses = {x["mes"]: x for x in r.json()["rendiciones"]}
    assert len(meses) == 6 and meses["2026-12"]["estado"] == "Aprobada"
    assert meses["2026-12"]["monto"] == 500_000

    # Acortar quitando meses vacíos sí se puede
    assert len(periodo(client, p["id"], "2026-10", "2026-12").json()["rendiciones"]) == 3


def test_actualizar_rendicion(client):
    p = nuevo(client)
    rid = periodo(client, p["id"], "2026-09", "2026-10").json()["rendiciones"][0]["id"]
    r = client.put(f"/api/rendiciones/{rid}", json={
        "estado": "Con observaciones", "monto": 1_200_000, "fecha_entrega": "2026-10-05",
        "fecha_revision": "2026-10-08", "observaciones": "  Faltan 3 boletas  ",
    })
    assert r.status_code == 200
    sep = r.json()["rendiciones"][0]
    assert sep["estado"] == "Con observaciones" and sep["observaciones"] == "Faltan 3 boletas"
    assert sep["fecha_entrega"] == "2026-10-05"
    assert r.json()["bitacora"][0]["texto"] == \
        "Rendición sep 2026: Pendiente → Con observaciones. Faltan 3 boletas"

    assert client.put(f"/api/rendiciones/{rid}", json={"estado": "Perdida"}).status_code == 422
    assert client.put(f"/api/rendiciones/{rid}", json={"monto": -5}).status_code == 422
    assert client.put("/api/rendiciones/9999", json={}).status_code == 404
    assert periodo(client, p["id"], "2026-12", "2026-09").status_code == 422
    assert periodo(client, p["id"], "2026-01", "2028-12").status_code == 422   # > 24 meses


def test_exportar_excel(client):
    from io import BytesIO
    from openpyxl import load_workbook

    p = nuevo(client, codigo="8%-CUL-001", monto=3_000_000)
    rend = periodo(client, p["id"], "2026-09", "2026-11").json()["rendiciones"]
    client.put(f"/api/rendiciones/{rend[0]['id']}", json={"estado": "Aprobada", "monto": 1_000_000})
    client.put(f"/api/rendiciones/{rend[1]['id']}", json={"estado": "Incompleta", "monto": 400_000,
                                                          "observaciones": "Falta factura"})

    r = client.get("/api/exportar/rendiciones.xlsx")
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Resumen", "Detalle", "Transferencias"]

    resumen = [[c.value for c in fila] for fila in wb["Resumen"].iter_rows()]
    assert resumen[0][:14] == ["Código", "Proyecto", "Organización", "Línea", "Etapa", "Año", "Período",
                               "Monto proyecto", "Total rendido", "Por rendir", "Vence pagaré",
                               "sep 2026", "oct 2026", "nov 2026"]
    assert resumen[1][6] == "sep 2026 a nov 2026 (3)"
    assert resumen[1][8:14] == [1_400_000, 1_600_000, None, "Aprobada", "Incompleta", "Pendiente"]

    detalle = [[c.value for c in fila] for fila in wb["Detalle"].iter_rows()]
    assert len(detalle) == 4
    assert detalle[2][4] == "Incompleta" and detalle[2][8] == "Falta factura"


def test_excel_incluye_proyectos_sin_rendiciones(client):
    from io import BytesIO
    from openpyxl import load_workbook

    nuevo(client, nombre="Sin rendiciones aún", monto=2_000_000)
    wb = load_workbook(BytesIO(client.get("/api/exportar/rendiciones.xlsx").content))
    resumen = [[c.value for c in fila] for fila in wb["Resumen"].iter_rows()]
    assert len(resumen) == 2
    assert resumen[1][1] == "Sin rendiciones aún" and resumen[1][6] == "Sin período definido"
    assert resumen[1][9] == 2_000_000          # todo por rendir
    detalle = [[c.value for c in fila] for fila in wb["Detalle"].iter_rows()]
    assert "Aún no hay rendiciones" in detalle[1][1]


def test_conexion_usable_desde_otro_hilo(tmp_path, monkeypatch):
    # FastAPI puede abrir la conexión en un hilo y usarla en otro dentro de la misma
    # petición; sin check_same_thread=False eso daba errores 500 intermitentes.
    import threading
    from app.db import conectar

    monkeypatch.setenv("CARTERA_DB", str(tmp_path / "hilos.db"))
    conn = conectar()
    errores = []

    def usar():
        try:
            conn.execute("SELECT 1").fetchone()
        except Exception as e:  # noqa: BLE001
            errores.append(e)

    t = threading.Thread(target=usar)
    t.start()
    t.join()
    conn.close()
    assert not errores


def test_pasos_por_etapa(client):
    cfg = client.get("/api/config").json()
    pasos = cfg["pasos_por_etapa"]
    # Toda etapa tiene sus pasos, salvo Rendición (se calcula con las rendiciones) y las que tienen checklist
    con_checklist = set(cfg["checklist_por_etapa"])
    assert con_checklist == {"Adjudicado", "Convenio"}
    assert set(pasos) == set(cfg["etapas"]) - {cfg["etapa_rendiciones"]} - con_checklist
    assert pasos["Ejecución"] == ["Hacer seguimiento a la ejecución"]
    assert pasos["Cerrado"] == []   # sin sugerencia


def test_anio_de_la_convocatoria(client):
    from datetime import date

    assert nuevo(client)["anio"] == date.today().year            # por defecto, el año actual
    p = nuevo(client, nombre="Del concurso anterior", anio=2025)
    assert p["anio"] == 2025
    datos = {k: p[k] for k in ["nombre", "codigo", "linea", "organizacion", "monto", "etapa",
                               "contacto", "accion", "fecha", "notas"]}
    assert client.put(f"/api/proyectos/{p['id']}", json={**datos, "anio": 2024}).json()["anio"] == 2024
    # El formulario ya no envía el año: editar no debe cambiarlo
    assert client.put(f"/api/proyectos/{p['id']}", json=datos).json()["anio"] == 2024
    assert client.post("/api/proyectos", json={"nombre": "x", "anio": 1999}).status_code == 422


def test_migracion_etapas_anteriores_pasan_a_adjudicado(tmp_path):
    """Una base con proyectos en Postulación/Evaluación queda en Adjudicado al abrirla."""
    import sqlite3
    from app import db

    ruta = tmp_path / "vieja.db"
    c = sqlite3.connect(ruta)
    for n, sql in enumerate(db.MIGRACIONES[:6], start=1):
        c.executescript(sql)
        c.execute(f"PRAGMA user_version = {n}")
    t = "2026-09-01T12:00:00+00:00"
    c.execute("INSERT INTO proyectos (nombre, etapa, accion, creado, actualizado) VALUES "
              "('A', 'Postulación', 'Revisar que la postulación esté completa', ?, ?)", (t, t))
    c.execute("INSERT INTO proyectos (nombre, etapa, accion, creado, actualizado) VALUES "
              "('B', 'Evaluación', 'Llamar al municipio', ?, ?)", (t, t))
    c.execute("INSERT INTO proyectos (nombre, etapa, accion, creado, actualizado) VALUES "
              "('C', 'Rendición', 'Revisar rendición', ?, ?)", (t, t))
    c.commit()
    c.close()

    c = db.conectar(ruta)
    filas = {r["nombre"]: dict(r) for r in c.execute("SELECT nombre, etapa, accion FROM proyectos")}
    assert filas["A"]["etapa"] == "Adjudicado" and filas["A"]["accion"] == "Preparar el convenio"
    assert filas["B"]["etapa"] == "Adjudicado" and filas["B"]["accion"] == "Llamar al municipio"   # lo escrito se respeta
    assert filas["C"]["etapa"] == "Rendición"
    notas = [r["texto"] for r in c.execute("SELECT texto FROM bitacora WHERE texto LIKE 'Etapa:%' ORDER BY id")]
    assert "Etapa: Postulación → Adjudicado (el seguimiento ahora empieza en Adjudicado)" in notas
    assert len(notas) == 2
    c.close()
