import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("CARTERA_DB", str(tmp_path / "test.db"))
    from app.main import app

    with TestClient(app) as c:
        yield c


def proyecto(client, **extra):
    r = client.post("/api/proyectos", json={"nombre": "Con cuotas", "monto": 10_000_001, **extra})
    assert r.status_code == 201, r.text
    return r.json()


def cantidad(client, pid, n):
    return client.put(f"/api/proyectos/{pid}/cuotas/cantidad", json={"cantidad": n})


def test_dos_cuotas_reparten_el_monto(client):
    p = proyecto(client, etapa="Transferencia")
    r = cantidad(client, p["id"], 2)
    assert r.status_code == 200
    cuotas = r.json()["cuotas"]
    assert [c["numero"] for c in cuotas] == [1, 2]
    assert sum(c["monto"] for c in cuotas) == 10_000_001 and cuotas[0]["monto"] == 5_000_000
    assert all(c["estado"] == "Programada" for c in cuotas)
    assert r.json()["bitacora"][0]["texto"] == "Transferencia en 2 cuotas"

    # La lista trae las cuotas para las alertas
    lista = {x["id"]: x for x in client.get("/api/proyectos").json()}
    assert len(lista[p["id"]]["cuotas"]) == 2


def test_transferir_cuota_queda_en_bitacora(client):
    p = proyecto(client)
    c1 = cantidad(client, p["id"], 2).json()["cuotas"][0]
    r = client.put(f"/api/cuotas/{c1['id']}", json={"monto": 5_000_000, "fecha_programada": "2026-08-01",
                                                    "estado": "Transferida", "observaciones": " Res. 123 "})
    assert r.status_code == 200
    c = r.json()["cuotas"][0]
    assert c["estado"] == "Transferida" and c["fecha_transferencia"] and c["observaciones"] == "Res. 123"
    assert r.json()["bitacora"][0]["texto"].startswith("1ª cuota transferida: $5.000.000")

    # No se puede quitar una cuota transferida
    assert cantidad(client, p["id"], 1).status_code == 200            # quita la 2ª (programada)
    r = client.put(f"/api/proyectos/{p['id']}/cuotas/cantidad", json={"cantidad": 1})
    assert len(r.json()["cuotas"]) == 1
    assert client.put(f"/api/cuotas/{c1['id']}", json={"estado": "Perdida"}).status_code == 422
    assert client.put("/api/cuotas/999", json={}).status_code == 404
    assert cantidad(client, p["id"], 0).status_code == 422
    assert cantidad(client, p["id"], 3).status_code == 422   # solo 1 o 2 cuotas


def test_quitar_cuota_transferida_no_se_permite(client):
    p = proyecto(client)
    cuotas = cantidad(client, p["id"], 2).json()["cuotas"]
    client.put(f"/api/cuotas/{cuotas[1]['id']}", json={"estado": "Transferida"})
    r = cantidad(client, p["id"], 1)
    assert r.status_code == 409 and "2ª" in r.json()["detail"]


def test_pasar_a_ejecucion_exige_primera_cuota(client):
    p = proyecto(client, etapa="Transferencia")
    client.post(f"/api/proyectos/{p['id']}/bitacora", json={"texto": "Se envió la solicitud de transferencia"})
    r = client.post(f"/api/proyectos/{p['id']}/avanzar")
    assert r.status_code == 409 and "1ª cuota" in r.json()["detail"]

    cuotas = cantidad(client, p["id"], 2).json()["cuotas"]
    client.put(f"/api/cuotas/{cuotas[0]['id']}", json={"monto": cuotas[0]["monto"], "estado": "Transferida"})
    r = client.post(f"/api/proyectos/{p['id']}/avanzar")
    assert r.status_code == 200 and r.json()["etapa"] == "Ejecución"   # la 2ª puede seguir pendiente


def test_nueva_etapa_transferencia_parte_definiendo_cuotas(client):
    p = proyecto(client, etapa="Adjudicado")
    client.post(f"/api/proyectos/{p['id']}/bitacora", json={"texto": "Convenio listo"})
    r = client.post(f"/api/proyectos/{p['id']}/avanzar")          # Adjudicado → Convenio
    client.post(f"/api/proyectos/{p['id']}/bitacora", json={"texto": "Firmado"})
    r = client.post(f"/api/proyectos/{p['id']}/avanzar")          # Convenio → Transferencia
    assert r.json()["etapa"] == "Transferencia" and r.json()["accion"] == "Definir las cuotas de transferencia"


def test_excel_trae_hoja_de_transferencias(client):
    from io import BytesIO
    from openpyxl import load_workbook

    p = proyecto(client, codigo="8%-X-1")
    cuotas = cantidad(client, p["id"], 2).json()["cuotas"]
    client.put(f"/api/cuotas/{cuotas[0]['id']}", json={"monto": 5_000_000, "estado": "Transferida"})
    wb = load_workbook(BytesIO(client.get("/api/exportar/rendiciones.xlsx").content))
    assert wb.sheetnames == ["Resumen", "Detalle", "Transferencias"]
    filas = [[c.value for c in f] for f in wb["Transferencias"].iter_rows()]
    assert filas[1][:4] == ["8%-X-1", "Con cuotas", "1ª", 5_000_000] and filas[1][5] == "Transferida"
    assert filas[2][2] == "2ª" and filas[2][5] == "Programada"
